"""
SalaryLens AI by CodeSec — 5-Sheet Excel Workbook Generator
"""
import os
import shutil
from datetime import datetime, timezone, date
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.salary_detection.filter import sanitize_excel_formula


class ExcelWorkbookGenerator:
    """
    Generates and updates deterministic, 5-sheet financial Excel workbooks:
    1. Summary
    2. Bank_Account_holder_info
    3. Bank_Txns
    4. Salary_Analysis_Txns
    5. Salary_Results
    """

    # Styling Palette
    HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Navy Blue
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    SUBHEADER_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    CANDIDATE_SUMMARY_FILL = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")  # Light Indigo
    ALT_ROW_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    BORDER_THIN = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    @classmethod
    def _apply_sheet_styling(cls, ws, max_col: int):
        """Applies auto column widths and freeze panes."""
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        for col in range(1, max_col + 1):
            col_letter = get_column_letter(col)
            max_len = 0
            for row in range(1, min(ws.max_row + 1, 100)):
                val = ws.cell(row=row, column=col).value
                if val is not None:
                    s_val = str(val)
                    if len(s_val) > max_len:
                        max_len = len(s_val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    @classmethod
    def generate_initial_workbook(
        cls,
        statement: Any,
        account_info: Dict[str, Any],
        transactions: List[Dict[str, Any]],
        filtering_stats: Dict[str, Any],
        output_path: Path,
        extraction_warnings: Optional[List[str]] = None,
    ) -> Path:
        """
        Builds the initial 5-sheet workbook after extraction and deterministic filtering.
        Salary_Results is initialized with an "Analysis in progress" state.
        """
        output_path.parent.mkdir(parents=True, exist_ok=True)
        tmp_path = output_path.with_suffix(".tmp")

        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # ----------------------------------------------------
        # Sheet 1: Summary
        # ----------------------------------------------------
        ws_summary = wb.create_sheet(title="Summary")
        cls._populate_summary_sheet(
            ws=ws_summary,
            statement=statement,
            filtering_stats=filtering_stats,
            salary_status="Pending / In Progress",
            ai_cost=Decimal("0.00"),
            ai_tokens=0,
            warnings=extraction_warnings or [],
        )

        # ----------------------------------------------------
        # Sheet 2: Bank_Account_holder_info
        # ----------------------------------------------------
        ws_acct = wb.create_sheet(title="Bank_Account_holder_info")
        cls._populate_account_info_sheet(ws=ws_acct, account_info=account_info, statement_id=str(statement.id))

        # ----------------------------------------------------
        # Sheet 3: Bank_Txns (Complete Transaction Dataset)
        # ----------------------------------------------------
        ws_bank_txns = wb.create_sheet(title="Bank_Txns")
        cls._populate_bank_txns_sheet(ws=ws_bank_txns, transactions=transactions)

        # ----------------------------------------------------
        # Sheet 4: Salary_Analysis_Txns (Filtered Dataset)
        # ----------------------------------------------------
        ws_salary_txns = wb.create_sheet(title="Salary_Analysis_Txns")
        cls._populate_salary_analysis_txns_sheet(ws=ws_salary_txns, transactions=transactions)

        # ----------------------------------------------------
        # Sheet 5: Salary_Results (Initial / Pending State)
        # ----------------------------------------------------
        ws_results = wb.create_sheet(title="Salary_Results")
        cls._populate_initial_results_sheet(ws=ws_results)

        wb.save(str(tmp_path))
        shutil.move(str(tmp_path), str(output_path))
        return output_path

    @classmethod
    def finalize_workbook(
        cls,
        workbook_path: Path,
        statement: Any,
        candidate_groups: List[Any],
        analysis_run: Optional[Any],
        transactions: List[Dict[str, Any]],
        filtering_stats: Dict[str, Any],
        ai_cost: Decimal = Decimal("0.00"),
        ai_tokens: int = 0,
        reviewer_decision: Optional[Dict[str, Any]] = None,
    ) -> Path:
        """
        Updates Sheet 1 (Summary) and Sheet 5 (Salary_Results) with validated AI and financial results.
        """
        tmp_path = workbook_path.with_suffix(".tmp")
        wb = openpyxl.load_workbook(str(workbook_path))

        # 1. Update Summary Sheet
        if "Summary" in wb.sheetnames:
            wb.remove(wb["Summary"])
        ws_summary = wb.create_sheet(title="Summary", index=0)
        status_label = analysis_run.salary_status if analysis_run else "Incomplete"
        if reviewer_decision:
            status_label = f"Human Review: {reviewer_decision.get('status', 'Approved').upper()}"

        cls._populate_summary_sheet(
            ws=ws_summary,
            statement=statement,
            filtering_stats=filtering_stats,
            salary_status=status_label,
            ai_cost=ai_cost,
            ai_tokens=ai_tokens,
            warnings=[],
            is_final=True,
        )

        # 2. Update Salary_Results Sheet
        if "Salary_Results" in wb.sheetnames:
            wb.remove(wb["Salary_Results"])
        ws_results = wb.create_sheet(title="Salary_Results")
        cls._populate_finalized_results_sheet(
            ws=ws_results,
            candidate_groups=candidate_groups,
            analysis_run=analysis_run,
            transactions=transactions,
            reviewer_decision=reviewer_decision,
        )

        wb.save(str(tmp_path))
        shutil.move(str(tmp_path), str(workbook_path))
        return workbook_path

    # =========================================================================
    # Sheet Population Helpers
    # =========================================================================

    @classmethod
    def _populate_summary_sheet(
        cls,
        ws,
        statement: Any,
        filtering_stats: Dict[str, Any],
        salary_status: str,
        ai_cost: Decimal,
        ai_tokens: int,
        warnings: List[str],
        is_final: bool = False,
    ):
        ws.title = "Summary"
        headers = ["Metric / Attribute", "Value"]
        ws.append(headers)
        for col_idx in [1, 2]:
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="left", vertical="center")

        total_txs = filtering_stats.get("total_transactions", 0)
        total_debits = filtering_stats.get("total_debits", 0)
        total_credits = filtering_stats.get("total_credits", 0)
        included_count = filtering_stats.get("included_in_salary_analysis", 0)
        excluded_count = filtering_stats.get("excluded_from_salary_analysis", 0)

        rows = [
            ("Statement ID", str(statement.id)),
            ("Original Filename", statement.filename),
            ("Processing Status", statement.status),
            ("Statement Start Date", str(statement.statement_start_date or "N/A")),
            ("Statement End Date", str(statement.statement_end_date or "N/A")),
            ("Total Statement Pages", statement.page_count),
            ("Pages Processed", statement.page_count),
            ("Pages Processed (Native PDF)", getattr(statement, "native_pages_count", statement.page_count)),
            ("Pages Processed (OCR Fallback)", getattr(statement, "ocr_pages_count", 0)),
            ("Total Extracted Transactions", total_txs),
            ("Total Debit Transactions", total_debits),
            ("Total Credit Transactions", total_credits),
            ("Transactions Included in Salary Analysis", included_count),
            ("Transactions Excluded from Salary Analysis", excluded_count),
        ]

        # Exclusion reasons breakdown
        exclusion_reasons = filtering_stats.get("exclusion_counts_by_reason", {})
        for reason, count in exclusion_reasons.items():
            rows.append((f"  └─ Excluded: {reason}", count))

        rows.extend([
            ("Potential Salary Transactions", included_count),
            ("Deterministic Validation Status", getattr(statement, "workbook_validation_status", "Verified")),
            ("Salary Analysis Status", salary_status),
            ("AI Tokens Consumed", ai_tokens),
            ("AI Total Cost (USD)", f"${float(ai_cost):.4f}"),
            ("Workbook Generation Timestamp", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")),
        ])

        if warnings:
            rows.append(("Extraction Warnings", "; ".join(warnings)))

        for r_idx, (k, v) in enumerate(rows, start=2):
            c1 = ws.cell(row=r_idx, column=1, value=sanitize_excel_formula(k))
            c2 = ws.cell(row=r_idx, column=2, value=sanitize_excel_formula(str(v)))
            c1.border = cls.BORDER_THIN
            c2.border = cls.BORDER_THIN
            if r_idx % 2 == 1:
                c1.fill = cls.ALT_ROW_FILL
                c2.fill = cls.ALT_ROW_FILL

        ws.column_dimensions["A"].width = 44
        ws.column_dimensions["B"].width = 48
        ws.views.sheetView[0].showGridLines = True

    @classmethod
    def _populate_account_info_sheet(cls, ws, account_info: Dict[str, Any], statement_id: str):
        headers = [
            "bank_data_customer_account_info_id",
            "start_date",
            "end_date",
            "bank_data_id",
            "name",
            "mobile",
            "landline",
            "address",
            "email",
            "pan",
        ]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.border = cls.BORDER_THIN

        row_vals = [
            sanitize_excel_formula(account_info.get("bank_data_customer_account_info_id", "")),
            sanitize_excel_formula(account_info.get("start_date", "")),
            sanitize_excel_formula(account_info.get("end_date", "")),
            sanitize_excel_formula(statement_id),
            sanitize_excel_formula(account_info.get("name", "")),
            sanitize_excel_formula(account_info.get("mobile", "")),
            sanitize_excel_formula(account_info.get("landline", "")),
            sanitize_excel_formula(account_info.get("address", "")),
            sanitize_excel_formula(account_info.get("email", "")),
            sanitize_excel_formula(account_info.get("pan", "")),
        ]
        ws.append(row_vals)
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=2, column=col_idx).border = cls.BORDER_THIN

        cls._apply_sheet_styling(ws, len(headers))

    @classmethod
    def _populate_bank_txns_sheet(cls, ws, transactions: List[Dict[str, Any]]):
        headers = [
            "transaction_id",
            "bank_data_transactions_info_id",
            "transaction_date",
            "value_date",
            "narration",
            "amount",
            "debit_amount",
            "credit_amount",
            "transaction_direction",
            "balance",
            "category",
            "chqno",
            "currency",
            "source_page",
            "extraction_method",
            "extraction_confidence",
            "validation_status",
            "validation_warnings",
            "included_in_salary_analysis",
            "exclusion_reason",
        ]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.border = cls.BORDER_THIN

        for r_idx, tx in enumerate(transactions, start=2):
            cr = tx.get("credit_amount")
            dr = tx.get("debit_amount")
            eff_amt = cr if cr is not None and cr > Decimal("0.00") else dr

            row = [
                sanitize_excel_formula(str(tx.get("id", ""))),
                sanitize_excel_formula(str(tx.get("id", ""))),  # stable ID mapped
                str(tx.get("transaction_date", "")),
                str(tx.get("value_date") or ""),
                sanitize_excel_formula(tx.get("raw_description") or tx.get("description", "")),
                float(eff_amt) if eff_amt is not None else "",
                float(dr) if dr is not None else "",
                float(cr) if cr is not None else "",
                tx.get("transaction_direction", ""),
                float(tx["balance"]) if tx.get("balance") is not None else "",
                sanitize_excel_formula(tx.get("category", "")),
                sanitize_excel_formula(tx.get("chqno", "")),
                tx.get("currency", "USD"),
                tx.get("page_number", 1),
                tx.get("extraction_method", "native"),
                float(tx["confidence"]) if tx.get("confidence") is not None else 1.0,
                tx.get("validation_status", "valid"),
                sanitize_excel_formula(tx.get("validation_warnings", "")),
                "TRUE" if tx.get("included_in_salary_analysis") else "FALSE",
                sanitize_excel_formula(tx.get("exclusion_reason", "")),
            ]
            ws.append(row)

            fill = cls.ALT_ROW_FILL if r_idx % 2 == 1 else None
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.border = cls.BORDER_THIN
                if fill:
                    cell.fill = fill
                # Format numeric currency columns
                if col_idx in (6, 7, 8, 10) and isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"

        cls._apply_sheet_styling(ws, len(headers))

    @classmethod
    def _populate_salary_analysis_txns_sheet(cls, ws, transactions: List[Dict[str, Any]]):
        headers = [
            "transaction_id",
            "transaction_date",
            "narration",
            "credit_amount",
            "transaction_direction",
            "employer_candidate",
            "currency",
            "source_page",
            "extraction_confidence",
            "validation_status",
            "inclusion_reason",
        ]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.border = cls.BORDER_THIN

        filtered_txs = [t for t in transactions if t.get("included_in_salary_analysis")]
        for r_idx, tx in enumerate(filtered_txs, start=2):
            cr = tx.get("credit_amount")
            row = [
                sanitize_excel_formula(str(tx.get("id", ""))),
                str(tx.get("transaction_date", "")),
                sanitize_excel_formula(tx.get("raw_description") or tx.get("description", "")),
                float(cr) if cr is not None else "",
                tx.get("transaction_direction", "CREDIT"),
                sanitize_excel_formula(tx.get("candidate_employer", "")),
                tx.get("currency", "USD"),
                tx.get("page_number", 1),
                float(tx["confidence"]) if tx.get("confidence") is not None else 1.0,
                tx.get("validation_status", "valid"),
                sanitize_excel_formula(tx.get("inclusion_reason", "")),
            ]
            ws.append(row)

            fill = cls.ALT_ROW_FILL if r_idx % 2 == 1 else None
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.border = cls.BORDER_THIN
                if fill:
                    cell.fill = fill
                if col_idx == 4 and isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"

        cls._apply_sheet_styling(ws, len(headers))

    @classmethod
    def _populate_initial_results_sheet(cls, ws):
        headers = [
            "salary_candidate_id",
            "employer_candidate",
            "supporting_transaction_id",
            "transaction_date",
            "narration",
            "salary_credit_amount",
            "currency",
            "suspected_salary_month",
            "candidate_status",
            "classification_confidence",
            "evidence_reason",
            "supporting_transaction_ids",
            "detected_frequency",
            "average_monthly_salary",
            "median_monthly_salary",
            "minimum_monthly_salary",
            "maximum_monthly_salary",
            "months_with_salary",
            "missing_months",
            "AI_explanation",
            "review_required",
            "review_reason",
            "validation_status",
        ]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.border = cls.BORDER_THIN

        # Placeholder row indicating pending analysis
        ws.append([
            "PENDING",
            "Salary detection in progress",
            "",
            "",
            "Awaiting OpenRouter Jev decision evaluation",
            "",
            "",
            "",
            "pending",
            0.0,
            "Initial workbook generated prior to AI classification",
            "",
            "",
            "",
            "",
            "",
            "",
            0,
            "",
            "Automated salary analysis is currently executing in background worker.",
            True,
            "Awaiting pipeline completion",
            "pending",
        ])
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=2, column=col_idx).border = cls.BORDER_THIN

        cls._apply_sheet_styling(ws, len(headers))

    @classmethod
    def _populate_finalized_results_sheet(
        cls,
        ws,
        candidate_groups: List[Any],
        analysis_run: Optional[Any],
        transactions: List[Dict[str, Any]],
        reviewer_decision: Optional[Dict[str, Any]] = None,
    ):
        headers = [
            "record_type",  # CANDIDATE_SUMMARY or SUPPORTING_TRANSACTION
            "salary_candidate_id",
            "employer_candidate",
            "supporting_transaction_id",
            "transaction_date",
            "narration",
            "salary_credit_amount",
            "currency",
            "suspected_salary_month",
            "candidate_status",
            "classification_confidence",
            "evidence_reason",
            "supporting_transaction_ids",
            "detected_frequency",
            "average_monthly_salary",
            "median_monthly_salary",
            "minimum_monthly_salary",
            "maximum_monthly_salary",
            "months_with_salary",
            "missing_months",
            "AI_explanation",
            "review_required",
            "review_reason",
            "validation_status",
        ]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.border = cls.BORDER_THIN

        tx_map = {str(t.get("id")): t for t in transactions}
        curr_row = 2

        # Format explanation
        ai_exp = ""
        if analysis_run:
            if isinstance(analysis_run.explanation, list):
                ai_exp = " | ".join(analysis_run.explanation)
            elif analysis_run.explanation:
                ai_exp = str(analysis_run.explanation)

        is_review_req = analysis_run.requires_human_review if analysis_run else True
        if reviewer_decision:
            is_review_req = False
            ai_exp = f"[Human Reviewer Confirmed: {reviewer_decision.get('review_notes', 'Approved')}] {ai_exp}"

        for group in candidate_groups:
            # 1. Summary Row for the Candidate Group
            grp_employer = getattr(group, "employer_name", "")
            if reviewer_decision and reviewer_decision.get("confirmed_employer"):
                grp_employer = reviewer_decision["confirmed_employer"]

            avg_amt = float(getattr(group, "average_amount", 0))
            med_amt = float(getattr(group, "median_amount", 0))
            min_amt = float(getattr(group, "min_amount", 0))
            max_amt = float(getattr(group, "max_amount", 0))
            months_count = getattr(group, "months_count", 0)
            missing_months = ", ".join(getattr(group, "missing_months", []) or [])

            # Extract supporting transaction IDs
            sup_tx_ids = []
            if hasattr(group, "transactions") and group.transactions:
                sup_tx_ids = [str(item.get("id")) for item in group.transactions if item.get("id")]
            elif hasattr(group, "candidates") and group.candidates:
                sup_tx_ids = [str(c.transaction_id) for c in group.candidates if c.transaction_id]

            summary_row = [
                "CANDIDATE_SUMMARY",
                str(getattr(group, "id", "")),
                sanitize_excel_formula(grp_employer),
                "",  # empty on candidate summary
                "",  # empty on candidate summary
                "",  # empty on candidate summary
                getattr(transactions[0], "currency", "USD") if transactions else "USD",
                "",  # empty on candidate summary
                analysis_run.salary_status if analysis_run else "evaluated",
                1.0,
                f"Consolidated Candidate Group ({len(sup_tx_ids)} credits across {months_count} months)",
                ", ".join(sup_tx_ids),
                "Monthly (Regular)" if getattr(group, "recurrence_score", 0) > Decimal("0.8") else "Variable",
                avg_amt,
                med_amt,
                min_amt,
                max_amt,
                months_count,
                missing_months,
                sanitize_excel_formula(ai_exp),
                "YES" if is_review_req else "NO",
                "; ".join(getattr(analysis_run, "reason_codes", []) or []) if analysis_run else "",
                "VALIDATED_BY_PYTHON_ENGINE",
            ]
            ws.append(summary_row)

            # Style Candidate Summary Row
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=curr_row, column=col_idx)
                c.fill = cls.CANDIDATE_SUMMARY_FILL
                c.font = Font(name="Calibri", size=10, bold=True)
                c.border = cls.BORDER_THIN
                if col_idx in (14, 15, 16, 17) and isinstance(c.value, (int, float)):
                    c.number_format = "#,##0.00"

            curr_row += 1

            # 2. Individual Supporting Transaction Records
            for tid in sup_tx_ids:
                tx_record = tx_map.get(tid)
                if not tx_record:
                    continue

                tx_date_str = str(tx_record.get("transaction_date", ""))
                month_label = tx_date_str[:7] if len(tx_date_str) >= 7 else ""
                cr_amt = tx_record.get("credit_amount")

                detail_row = [
                    "SUPPORTING_TRANSACTION",
                    str(getattr(group, "id", "")),
                    sanitize_excel_formula(grp_employer),
                    sanitize_excel_formula(tid),
                    tx_date_str,
                    sanitize_excel_formula(tx_record.get("raw_description") or tx_record.get("description", "")),
                    float(cr_amt) if cr_amt is not None else 0.0,
                    tx_record.get("currency", "USD"),
                    month_label,
                    "confirmed_credit",
                    float(tx_record.get("confidence") or 1.0),
                    f"Direct salary credit supporting {grp_employer}",
                    tid,
                    "",  # individual transaction does not re-duplicate frequency
                    "",  # individual transaction does not re-duplicate average
                    "",  # individual transaction does not re-duplicate median
                    "",  # individual transaction does not re-duplicate min
                    "",  # individual transaction does not re-duplicate max
                    "",  # individual transaction does not re-duplicate months
                    "",
                    "",
                    "",
                    "",
                    "VERIFIED_TRANSACTION",
                ]
                ws.append(detail_row)

                for col_idx in range(1, len(headers) + 1):
                    c = ws.cell(row=curr_row, column=col_idx)
                    c.border = cls.BORDER_THIN
                    if col_idx == 7 and isinstance(c.value, (int, float)):
                        c.number_format = "#,##0.00"

                curr_row += 1

        cls._apply_sheet_styling(ws, len(headers))
