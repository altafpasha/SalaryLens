"""
SalaryLens AI by CodeSec — Comprehensive Tests for 5-Sheet Excel Workbook and Secure Download API
"""
import os
import uuid
import openpyxl
import pytest
import httpx
from datetime import datetime, timezone, date
from decimal import Decimal
from pathlib import Path

from app.models.statement import Statement
from app.models.transaction import Transaction
from app.services.reporting.workbook import ExcelWorkbookGenerator
from app.services.document_processing.account_info import AccountInfoExtractor
from app.services.salary_detection.filter import TransactionFilterService


def test_account_info_extractor_deterministic_and_no_hallucination():
    sample_text = """
    GLOBAL METRO BANK — ACCOUNT STATEMENT
    Customer Name: Alex Morgan
    Account No: 8849-2910-4491
    Statement Period: 2026-03-01 to 2026-05-31
    Email: alex.morgan@example.com
    Phone No: +1-555-0199
    Communication Address: 742 Evergreen Terrace, Springfield, OR 97477
    """
    info = AccountInfoExtractor.extract_from_text(sample_text)

    assert info["name"] == "Alex Morgan"
    assert info["account_number"] == "8849-2910-4491"
    assert info["start_date"] == "2026-03-01"
    assert info["end_date"] == "2026-05-31"
    assert info["email"] == "alex.morgan@example.com"
    assert "+1-555-0199" in info["mobile"]
    assert "742 Evergreen Terrace" in info["address"]
    # Unextractable fields must remain blank, never fabricated
    assert info["pan"] == ""
    assert info["landline"] == ""


def test_generate_initial_and_finalized_5_sheet_workbook(tmp_path: Path):
    stmt_id = uuid.uuid4()
    org_id = uuid.uuid4()

    statement_mock = type("MockStatement", (), {
        "id": stmt_id,
        "filename": "april_may_statement.pdf",
        "status": "completed",
        "page_count": 5,
        "native_pages_count": 5,
        "ocr_pages_count": 0,
        "statement_start_date": date(2026, 4, 1),
        "statement_end_date": date(2026, 5, 31),
        "workbook_validation_status": "Verified",
    })()

    account_info = {
        "bank_data_customer_account_info_id": str(uuid.uuid4()),
        "start_date": "2026-04-01",
        "end_date": "2026-05-31",
        "name": "Jane Doe",
        "mobile": "9876543210",
        "email": "jane@doe.com",
        "address": "123 Main St",
        "pan": "",
        "landline": "",
    }

    raw_txs = [
        {
            "id": "tx-1",
            "transaction_date": "2026-04-30",
            "value_date": "2026-04-30",
            "raw_description": "=HYPERLINK('http://malicious.com', 'SALARY')",  # Injection test!
            "credit_amount": Decimal("5200.00"),
            "debit_amount": None,
            "balance": Decimal("10000.00"),
            "currency": "USD",
            "page_number": 1,
            "is_salary_candidate": True,
            "candidate_employer": "ACME TECH",
            "confidence": 0.95,
        },
        {
            "id": "tx-2",
            "transaction_date": "2026-05-02",
            "value_date": "2026-05-02",
            "raw_description": "-ATM CASH WDL #55",  # Injection test!
            "credit_amount": None,
            "debit_amount": Decimal("200.00"),
            "balance": Decimal("9800.00"),
            "currency": "USD",
            "page_number": 1,
            "is_salary_candidate": False,
            "candidate_employer": None,
            "confidence": 0.0,
        },
    ]

    annotated_txs, filtering_stats = TransactionFilterService.filter_statement_transactions(raw_txs)
    output_xlsx = tmp_path / "test_statement.xlsx"

    # 1. Generate Initial Workbook
    ExcelWorkbookGenerator.generate_initial_workbook(
        statement=statement_mock,
        account_info=account_info,
        transactions=annotated_txs,
        filtering_stats=filtering_stats,
        output_path=output_xlsx,
    )

    assert output_xlsx.exists()
    wb = openpyxl.load_workbook(str(output_xlsx))

    # All 5 worksheets must be present
    expected_sheets = ["Summary", "Bank_Account_holder_info", "Bank_Txns", "Salary_Analysis_Txns", "Salary_Results"]
    assert wb.sheetnames == expected_sheets

    # Verify Sheet 1: Summary contents
    ws_sum = wb["Summary"]
    assert ws_sum.cell(row=1, column=1).value == "Metric / Attribute"
    assert ws_sum.cell(row=2, column=2).value == str(stmt_id)

    # Verify Sheet 2: Account info
    ws_acct = wb["Bank_Account_holder_info"]
    assert ws_acct.cell(row=2, column=5).value == "Jane Doe"
    assert ws_acct.cell(row=2, column=10).value is None or ws_acct.cell(row=2, column=10).value == ""

    # Verify Sheet 3: Bank_Txns contains ALL transactions and escaped formula
    ws_txns = wb["Bank_Txns"]
    assert ws_txns.max_row == 3  # Header + 2 txs
    escaped_narration = ws_txns.cell(row=2, column=5).value
    assert escaped_narration.startswith("'="), "Formula injection not escaped!"

    # Verify Sheet 4: Salary_Analysis_Txns contains only included transactions
    ws_sal_tx = wb["Salary_Analysis_Txns"]
    assert ws_sal_tx.max_row == 2  # Header + 1 included salary credit

    # 2. Finalize Workbook with Candidate Groups and Validated AI Results
    candidate_group_mock = type("MockCandidateGroup", (), {
        "id": "cg-1",
        "employer_name": "ACME TECH",
        "average_amount": Decimal("5200.00"),
        "median_amount": Decimal("5200.00"),
        "min_amount": Decimal("5200.00"),
        "max_amount": Decimal("5200.00"),
        "months_count": 1,
        "missing_months": [],
        "recurrence_score": Decimal("0.95"),
        "transactions": [{"id": "tx-1"}],
    })()

    analysis_run_mock = type("MockAnalysisRun", (), {
        "salary_status": "likely_salary",
        "explanation": "Verified continuous payroll credits from ACME TECH.",
        "requires_human_review": False,
        "reason_codes": [],
    })()

    ExcelWorkbookGenerator.finalize_workbook(
        workbook_path=output_xlsx,
        statement=statement_mock,
        candidate_groups=[candidate_group_mock],
        analysis_run=analysis_run_mock,
        transactions=annotated_txs,
        filtering_stats=filtering_stats,
        ai_cost=Decimal("0.0250"),
        ai_tokens=1500,
    )

    wb_final = openpyxl.load_workbook(str(output_xlsx))
    assert wb_final.sheetnames == expected_sheets

    # Verify Sheet 5: Salary_Results populated
    ws_res = wb_final["Salary_Results"]
    # Row 2: Candidate summary
    assert ws_res.cell(row=2, column=1).value == "CANDIDATE_SUMMARY"
    assert ws_res.cell(row=2, column=3).value == "ACME TECH"
    assert ws_res.cell(row=2, column=14).value == 5200.00

    # Row 3: Supporting individual transaction
    assert ws_res.cell(row=3, column=1).value == "SUPPORTING_TRANSACTION"
    assert ws_res.cell(row=3, column=4).value == "tx-1"
    assert ws_res.cell(row=3, column=7).value == 5200.00


@pytest.mark.asyncio
async def test_download_endpoint_authorization_and_isolation(client: httpx.AsyncClient):
    # 1. Register Tenant A
    ts = int(datetime.now().timestamp())
    reg_a = await client.post("/api/v1/auth/register", json={
        "email": f"auditor_a_{ts}@banka.com",
        "password": "SecurePassword123!",
        "full_name": "Auditor A",
        "organization_name": f"Bank A {ts}",
        "role": "reviewer",
    })
    assert reg_a.status_code == 201
    token_a = reg_a.json()["access_token"]
    headers_a = {"Authorization": f"Bearer {token_a}"}

    # 2. Register Tenant B
    reg_b = await client.post("/api/v1/auth/register", json={
        "email": f"auditor_b_{ts}@bankb.com",
        "password": "SecurePassword123!",
        "full_name": "Auditor B",
        "organization_name": f"Bank B {ts}",
        "role": "reviewer",
    })
    assert reg_b.status_code == 201
    token_b = reg_b.json()["access_token"]
    headers_b = {"Authorization": f"Bearer {token_b}"}

    # 3. Tenant A uploads a statement PDF
    with open("fixtures/synthetic_statements/clean_text_statement.pdf", "rb") as f:
        pdf_bytes = f.read()

    up_resp = await client.post(
        "/api/v1/statements",
        files={"file": ("stmt_a.pdf", pdf_bytes, "application/pdf")},
        data={"duration_months": "3"},
        headers=headers_a,
    )
    assert up_resp.status_code == 202
    stmt_id = up_resp.json()["statement_id"]

    # 4. Trigger processing pipeline synchronously
    from app.workers.tasks import run_statement_processing_pipeline
    org_id_a = reg_a.json()["organization_id"]
    success = await run_statement_processing_pipeline(str(stmt_id), str(org_id_a))
    assert success is True

    # 5. Check status endpoint contains workbook metadata
    status_resp = await client.get(f"/api/v1/statements/{stmt_id}/status", headers=headers_a)
    assert status_resp.status_code == 200
    st_data = status_resp.json()
    assert st_data["workbook_status"] in ["initial", "completed"]
    assert st_data["total_transactions_count"] > 0
    assert st_data["included_in_salary_count"] > 0

    # 6. Tenant A downloads the workbook (Authorized)
    dl_resp = await client.get(f"/api/v1/statements/{stmt_id}/download", headers=headers_a)
    assert dl_resp.status_code == 200
    assert dl_resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "SalaryLens_" in dl_resp.headers.get("content-disposition", "")
    assert len(dl_resp.content) > 1000

    # 7. Tenant B attempts to download Tenant A's workbook (Cross-Tenant Blocked)
    cross_dl = await client.get(f"/api/v1/statements/{stmt_id}/download", headers=headers_b)
    assert cross_dl.status_code == 404

    # 8. Non-existent statement download attempt
    fake_id = str(uuid.uuid4())
    fake_dl = await client.get(f"/api/v1/statements/{fake_id}/download", headers=headers_a)
    assert fake_dl.status_code == 404
